#!/usr/bin/env python3
"""Night Plant themed 10-slide deck: Future of Smart Buildings."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x0B, 0x12, 0x20)
PANEL = RGBColor(0x12, 0x1A, 0x2C)
LIVE = RGBColor(0x3D, 0xDC, 0x97)
GOLD = RGBColor(0xE8, 0xC5, 0x47)
AIR = RGBColor(0x7E, 0xB6, 0xFF)
TYPE = RGBColor(0xF4, 0xF1, 0xE8)
MUTE = RGBColor(0x8B, 0x93, 0xA7)
RULE = RGBColor(0x2A, 0x35, 0x4D)

W = Inches(13.333)
H = Inches(7.5)
OUT = Path("/Users/nithya-verity/Downloads/Nava Studios/Future-of-Smart-Buildings-BMS-IoT-AI.pptx")

SLIDES = [
    {
        "kicker": "01  /  OPENING",
        "title": "The building is now a machine",
        "headline": "Buildings eat 30% of the world’s energy. Most of them still run on guesswork.",
        "bullets": [
            "IEA: buildings take ~30% of global final energy and ~26% of energy-related CO2.",
            "UNEP: buildings & construction ~37% of energy-related CO2 — retrofit decides the next 20 years.",
            "Most commercial HVAC still sequences on clocks and design loads, not occupancy, weather, or tariff.",
            "The bottleneck is integration and control logic — not another chiller brand.",
            "Mech owns the physics; EEE owns the electrons. Software without that literacy fails in the plant room.",
        ],
        "visual": "Visual: night tower + ducts/busbars. One huge 30%.",
        "note": "Open with the 10-second hook, then land 30% on screen. Pause after “guesswork.” This hour is a map of where Mechanical and EEE work sits in a smart building — not a vendor tour.",
        "stat": "30%",
        "stat_l": "of global final energy\nin buildings (IEA)",
    },
    {
        "kicker": "02  /  DEFINITION",
        "title": "What “smart” actually means",
        "headline": "A smart building is a closed-loop control system with a carbon and rupee constraint.",
        "bullets": [
            "Sense: T, RH, CO2, occupancy, kW/kWh, flow, pressure, lux, equipment status.",
            "Decide: BMS sequences and setpoints — ASHRAE logic before apps.",
            "Actuate: valves, dampers, VFDs, lighting, ATS/DG, chillers, pumps.",
            "Verify: submetering and trends. If you cannot prove 12%, you did not save 12%.",
            "Example: occupancy + CO2 resets outdoor air; tariff sheds non-critical loads.",
        ],
        "visual": "Visual: four-layer stack with Mech vs EEE ownership tags.",
        "note": "Draw the loop if you can. Ask: where does your final-year project sit — sensor, controller, or actuator?",
        "stat": "S→D→A→V",
        "stat_l": "Sense · Decide\nActuate · Verify",
    },
    {
        "kicker": "03  /  BMS",
        "title": "The building’s operating system",
        "headline": "If the BMS is weak, IoT and AI are decoration on a broken plant.",
        "bullets": [
            "A BMS is a distributed control system: controllers, I/O, workstation, trends, alarms — cousin to DCS/SCADA.",
            "Sequences (ASHRAE Guideline 36 trim-and-respond) can cut fan energy 10–30% vs fixed setpoint.",
            "Mid-size office: often 2,000–8,000 points. Naming and calibration decide maintainability.",
            "Meters, DG, UPS, lifts, lighting on one time-synced historian.",
            "Commissioning / retro-commissioning often finds 5–20% before any AI purchase.",
        ],
        "visual": "Visual: chiller → pumps → AHUs → VAV with BACnet objects.",
        "note": "BMS work is control engineering, not clicking graphics. 24×7 IT campuses: latent load and unreliable grid make sequences and backup power the same design.",
        "stat": "5–20%",
        "stat_l": "typical savings from\ncommissioning alone",
    },
    {
        "kicker": "04  /  IOT",
        "title": "From sparse sensors to a dense twin",
        "headline": "IoT does not replace the BMS. It densifies the measurement layer it always needed.",
        "bullets": [
            "Classic BMS: a few zone sensors. IoT: occupancy, IAQ, plug meters, vibration — often 5–10× more points.",
            "Wireless (LoRaWAN, BLE, Wi-Fi, NB-IoT) instruments existing campuses without tearing every ceiling.",
            "Occupancy-driven HVAC commonly shows 10–25% HVAC reduction when schedules match reality.",
            "EEE: CT/PT accuracy and OT security. A 2% CT error wrecks a 5% “AI savings” claim.",
            "Mech: sensor placement, stratification, dirty filters that make a dashboard lie.",
        ],
        "visual": "Visual: occupancy heatmap vs a dirty, uncalibrated duct sensor.",
        "note": "IoT without calibration is noise. Story: a CO2 sensor in return air that “proved” a packed exam hall was empty.",
        "stat": "10–25%",
        "stat_l": "HVAC energy when\nschedules match occupancy",
    },
    {
        "kicker": "05  /  AI",
        "title": "From rules to prediction",
        "headline": "Use AI to find the broken valve and tomorrow’s peak — not to reinvent thermodynamics.",
        "bullets": [
            "Fault detection (FDD): stuck dampers, simultaneous heat/cool, fouled condensers — 10–30% avoidable waste.",
            "Load and occupancy forecasting for chiller staging, storage, and DG/grid decisions.",
            "Supervisory optimization proposes setpoints; BMS still executes interlocks, freeze protection, fire modes.",
            "Data-centre ML cooling is a cousin idea — occupied buildings are messier. Transfer the method, not the hype.",
            "Guardrail: a model that “saves energy” by starving outdoor air is a health failure.",
        ],
        "visual": "Visual: FDD fault tree | 24-hour kW forecast vs demand-limit band.",
        "note": "Say slowly: AI is a supervisor, BMS is the operator, physics is the law. Example: detecting a 3°C chilled-water delta-T collapse before the bill arrives.",
        "stat": "AI ⊂ BMS",
        "stat_l": "supervisor, not\nthe safety layer",
    },
    {
        "kicker": "06  /  ENERGY",
        "title": "The scoreboard and the market",
        "headline": "Energy management is not a green poster. It is metered, tariff-aware, plant-level control.",
        "bullets": [
            "Split meters: incoming HT/LT, chiller plant, lighting, plug, critical IT — or “8% savings” is a story.",
            "India: ToD tariffs, demand charges, DG at ~₹25–40/kWh. Shaving 100 kW at 18:00 can beat a 5% all-day kWh cut.",
            "Thermal storage, pre-cooling, and kW/TR are Mechanical core — now visible in real time.",
            "Solar + BMS interlocks so PV does not fight the diesel plant (EEE: inverters, harmonics, islanding).",
            "A plant at 0.55–0.70 kW/TR vs neglected 0.90–1.1 kW/TR is a larger lever than swapping every LED twice.",
        ],
        "visual": "Visual: Sankey — grid + DG + solar → chillers, AHUs, lighting, IT.",
        "note": "Put simple bill math on screen: demand vs energy charge. Tie it to VFDs, staging, and power factor.",
        "stat": "kW/TR",
        "stat_l": "the plant KPI that\nbeats a green poster",
    },
    {
        "kicker": "07  /  STACK",
        "title": "One building, four layers",
        "headline": "The future building is not four products. It is one control and data architecture.",
        "bullets": [
            "Working stack: calibrated sensors, documented sequences, open protocols, historian, KPIs operators actually open.",
            "Typical stack: ~10% from sequences, 5–15% from sensors/ops, plus peak cuts.",
            "Digital twins are useful when anchored to live points and physics — not pretty BIM with fake occupancy.",
            "Unsegmented BACnet on the corporate LAN is a design defect. OT/IT zoning belongs in EEE for buildings too.",
            "Operators override “because 4th floor is always hot.” Fix the floor, or the AI is dead by Friday.",
        ],
        "visual": "Visual: field → controllers → BMS/historian → analytics, with a security band across OT.",
        "note": "Systems-thinking slide. Point at Mech and EEE: who owns cybersecurity? Both, plus IT.",
        "stat": "OT ≠ IT",
        "stat_l": "segment the network\nor inherit the incident",
    },
    {
        "kicker": "08  /  CAREERS",
        "title": "Where Mech and EEE actually work",
        "headline": "The scarce graduate is not “someone who heard of IoT.” It is someone who can commission a plant.",
        "bullets": [
            "Mechanical: HVAC design, energy models, retro-commissioning, IAQ, fire-smoke control interfaces.",
            "EEE: ELV/BMS, SCADA, metering, VFD harmonics, protection, microgrids, EV charging load.",
            "Shared: controls engineer, energy manager, commissioning authority, campus ESG data owner.",
            "India demand: GCCs, hospitals, airports, data centres, manufacturing parks — plant + digital, not smart-city slogans.",
            "Hire-ready project: BACnet/Modbus logger + chiller kW/TR dashboard + a documented sequence.",
        ],
        "visual": "Visual: two-column job map; overlap = Controls & Energy.",
        "note": "Internships: FM, controls OEMs, ESCOs, data-centre MEP. Portfolio = points list + sequence + trend, not only CAD.",
        "stat": "Cx",
        "stat_l": "commissioning is\nthe bilingual skill",
    },
    {
        "kicker": "09  /  LIMITS",
        "title": "What can go wrong",
        "headline": "A building that saves energy by making people sick is not sustainable. It is a failed controller.",
        "bullets": [
            "Comfort, IAQ, and safety override energy KPIs — always. OA cutbacks without CO2/PM are a health failure.",
            "A ₹2 crore analytics overlay on failed valves is malpractice. 5–20% often sits in maintenance.",
            "Proprietary protocols and cloud-only historians create stranded assets.",
            "Rebound: report absolute kWh and intensity (kWh/m², kWh/occupant), not only % vs last year.",
            "Match compute to the plant. Edge FDD is cheap; giant remote models on a 200 TR plant can be theatre.",
        ],
        "visual": "Visual: risk matrix — IAQ, cyber, lock-in, payback, data quality.",
        "note": "This slide earns trust. Do not rush. You will use it when someone says AI is hype or smart buildings are only for rich campuses.",
        "stat": "IAQ > kWh",
        "stat_l": "energy never outranks\nhealth or fire mode",
    },
    {
        "kicker": "10  /  CLOSE",
        "title": "Make this one tell the truth",
        "headline": "Don’t wait for a smarter building. Learn to make this one tell the truth.",
        "bullets": [
            "Truth = calibrated sensors + documented sequences + submetering + an operator allowed to act.",
            "CTA: pick one AHU or one incoming feeder. Trend it 14 days. Bring waste in kWh, kW, or °C — not slogans.",
            "A commissioning walk-through teaches more than a week of smart-city videos.",
            "Build bilingual skill: one plant diagram + one electrical single-line + a points list.",
            "The industry does not need more people who can say “digital twin.” It needs people who can explain 1.0 kW/TR.",
        ],
        "visual": "Visual: dark slide, one sentence, one CTA. No collage.",
        "note": "Deliver the close standing still. End on the CTA, then the thank-you. Leave 8–10 minutes for hard questions. Thank them as future operators of India’s next 20 years of floor space.",
        "stat": "14 days",
        "stat_l": "one real loop,\nthen you have evidence",
    },
]


def set_run_font(run, name, size, color, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rPr.append(ea)
    ea.set("typeface", name)
    latin = rPr.find(qn("a:latin"))
    if ea is not None:
        pass
    if latin is None:
        latin = OxmlElement("a:latin")
        rPr.insert(0, latin)
    latin.set("typeface", name)


def add_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text
    for p in notes.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.name = "Calibri"


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(sh, color)
    return sh


def add_text(slide, l, t, w, h, text, name, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    set_run_font(run, name, size, color, bold)
    return box


def add_bullets(slide, l, t, w, h, items):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        # accent dash
        run = p.add_run()
        run.text = "▸  "
        set_run_font(run, "Calibri", 15, LIVE, True)
        run = p.add_run()
        run.text = item
        set_run_font(run, "Calibri", 15, TYPE, False)
    return box


def chrome(slide, kicker, title, n):
    add_rect(slide, 0, 0, W, H, INK)
    add_rect(slide, 0, 0, Inches(0.12), H, LIVE)
    add_text(slide, Inches(0.55), Inches(0.22), Inches(8.5), Inches(0.32), kicker, "Calibri", 11, LIVE, True)
    add_text(slide, Inches(0.55), Inches(0.48), Inches(9.2), Inches(0.55), title, "Calibri", 26, TYPE, True)
    add_rect(slide, Inches(0.55), Inches(1.08), Inches(2.1), Inches(0.045), GOLD)
    add_text(
        slide,
        Inches(10.4),
        Inches(0.22),
        Inches(2.5),
        Inches(0.28),
        "MECH  ×  EEE",
        "Calibri",
        11,
        MUTE,
        True,
        PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        Inches(10.4),
        Inches(7.12),
        Inches(2.5),
        Inches(0.25),
        f"{n:02d}  /  10",
        "Calibri",
        11,
        MUTE,
        False,
        PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        Inches(0.55),
        Inches(7.12),
        Inches(8.5),
        Inches(0.25),
        "Future of Smart Buildings  ·  BMS · IoT · AI · Energy Management",
        "Calibri",
        11,
        MUTE,
        False,
    )


def stat_panel(slide, stat, label):
    add_rect(slide, Inches(10.15), Inches(1.35), Inches(2.7), Inches(5.45), PANEL)
    add_rect(slide, Inches(10.15), Inches(1.35), Inches(0.08), Inches(5.45), GOLD)
    add_text(slide, Inches(10.35), Inches(1.65), Inches(2.35), Inches(0.28), "HOLD THIS", "Calibri", 10, GOLD, True)
    add_text(slide, Inches(10.35), Inches(2.15), Inches(2.35), Inches(1.35), stat, "Calibri", 28, LIVE, True)
    add_text(slide, Inches(10.35), Inches(3.55), Inches(2.35), Inches(1.4), label, "Calibri", 13, TYPE, False)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    for i, s in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        chrome(slide, s["kicker"], s["title"], i)
        add_text(
            slide,
            Inches(0.55),
            Inches(1.28),
            Inches(9.3),
            Inches(1.15),
            s["headline"],
            "Calibri",
            20,
            AIR,
            True,
        )
        add_bullets(slide, Inches(0.5), Inches(2.5), Inches(9.4), Inches(4.4), s["bullets"])
        add_text(slide, Inches(0.55), Inches(6.72), Inches(9.4), Inches(0.32), s["visual"], "Calibri", 11, MUTE, False)
        stat_panel(slide, s["stat"], s["stat_l"])
        add_notes(slide, s["note"])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
